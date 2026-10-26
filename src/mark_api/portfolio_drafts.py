"""Explicit local CSV/Excel intake into a separate, unpublished product-draft table.

Not an owner inventory; it never writes to ad_snapshots, authorizes a platform
operation, contacts a browser/AI provider, or makes SKU an ads platform ID.
All rows are validated before the single SQLite transaction. Import is replay-
idempotent and conflicts fail the entire batch without mutating existing rows.
"""
from __future__ import annotations

import csv
from contextlib import closing
from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import secrets
import sqlite3
import stat
from typing import Any
import zipfile

from .storage import SnapshotStore

MAX_ROWS = 5000
MAX_FILE_BYTES = 10 * 1024 * 1024
_MAX_DECOMPRESSED = 36 * 1024 * 1024
_MAX_CELLS = 7
_FIELDS = ("sku", "title", "description", "category", "price_eur", "images")
_REQUIRED = _FIELDS[:-1]
_ALIASES = {
    "sku": ("sku", "artikelnummer", "artikel_nr", "artikel-nr", "artnr", "product_id"),
    "title": ("title", "titel", "bezeichnung", "artikelname"),
    "description": ("description", "beschreibung", "text"),
    "category": ("category", "kategorie", "kategoriepfad"),
    "price_eur": ("price_eur", "preis", "preis_eur", "festpreis"),
    "images": ("images", "bilder", "fotos", "bilddateien"),
}
_SKU = re.compile(r"[A-Za-z0-9][A-Za-z0-9._-]{0,79}\Z")
_IMG = re.compile(r"[A-Za-z0-9ÄÖÜäöüß][A-Za-z0-9ÄÖÜäöüß ._()-]{0,180}\.(?:jpg|jpeg|png|webp)\Z", re.IGNORECASE)
_HEX = re.compile(r"[0-9a-f]{64}\Z")
_DDL = """
CREATE TABLE IF NOT EXISTS portfolio_drafts (
  sku TEXT PRIMARY KEY COLLATE BINARY NOT NULL,
  title TEXT NOT NULL,
  description TEXT NOT NULL,
  category_json TEXT NOT NULL,
  price_eur INTEGER NOT NULL,
  images_json TEXT NOT NULL,
  source_sha256 TEXT NOT NULL,
  source_name TEXT NOT NULL,
  created_at TEXT NOT NULL,
  status TEXT NOT NULL DEFAULT 'draft' CHECK(status = 'draft')
)
"""


class DraftImportError(ValueError):
    def __init__(self, message: str, *, row: int | None = None, field: str = "") -> None:
        super().__init__(message)
        self.row = row
        self.field = field

    def as_issue(self) -> dict[str, object]:
        return {"row": self.row, "field": self.field, "message": str(self)}


class DraftBatchError(DraftImportError):
    def __init__(self, issues: list[dict[str, object]]) -> None:
        super().__init__("Importdatei enthaelt ungueltige Artikel.")
        self.issues = issues[:25]


@dataclass(frozen=True, slots=True)
class ProductDraft:
    sku: str
    title: str
    description: str
    category: tuple[str, ...]
    price_eur: int
    images: tuple[str, ...]

    def fields(self) -> tuple[object, ...]:
        return (self.title, self.description, json.dumps(self.category, ensure_ascii=True),
                self.price_eur, json.dumps(self.images, ensure_ascii=True))

    def as_dict(self) -> dict[str, object]:
        return {"sku": self.sku, "title": self.title, "description": self.description,
                "category": " > ".join(self.category), "price_eur": self.price_eur,
                "images": list(self.images), "status": "draft", "published": False}


def _read_file(path: Path) -> tuple[bytes, str]:
    if path.suffix.casefold() not in (".csv", ".xlsx"):
        raise DraftImportError("Nur CSV- oder XLSX-Dateien sind erlaubt.", field="file")
    if not path.is_absolute():
        raise DraftImportError("Dateipfad muss absolut sein.", field="file")
    try:
        meta = path.lstat()
        reparse = getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400)
        if not stat.S_ISREG(meta.st_mode) or (getattr(meta, "st_file_attributes", 0) & reparse):
            raise DraftImportError("Nur regulaere Dateien ohne Verknuepfung.", field="file")
        if not 0 < meta.st_size <= MAX_FILE_BYTES:
            raise DraftImportError("Dateigroesse unzulaessig (max. 10 MB).", field="file")
        flags = os.O_RDONLY | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0)
        fd = os.open(path, flags)
        with os.fdopen(fd, "rb") as stream:
            opened = os.fstat(stream.fileno())
            if not stat.S_ISREG(opened.st_mode) or opened.st_size != meta.st_size:
                raise DraftImportError("Datei hat sich beim Oeffnen geaendert.", field="file")
            content = stream.read(MAX_FILE_BYTES + 1)
            if len(content) != opened.st_size:
                raise DraftImportError("Datei hat sich beim Lesen geaendert.", field="file")
    except DraftImportError:
        raise
    except OSError:
        raise DraftImportError("Datei konnte nicht gelesen werden.", field="file") from None
    return content, hashlib.sha256(content).hexdigest()


def _column_map(headers: list[object]) -> dict[str, int]:
    if len(headers) > 40:
        raise DraftImportError("Zu viele Spalten (max. 40).", row=1)
    found: dict[str, int] = {}
    for i, raw in enumerate(headers):
        if not isinstance(raw, str):
            raise DraftImportError("Ungueltige Spaltenbezeichnung.", row=1)
        key = raw.strip().casefold().replace(" ", "_")
        matched = [field for field, aliases in _ALIASES.items() if key in aliases]
        if matched:
            field = matched[0]
            if field in found:
                raise DraftImportError("Doppelte Spalte: " + field, row=1, field=field)
            found[field] = i
        elif key:
            raise DraftImportError("Unbekannte Spalte: " + key[:50], row=1)
    for field in _REQUIRED:
        if field not in found:
            raise DraftImportError("Pflichtspalte fehlt: " + field, row=1, field=field)
    return found


def _text(value: object, field: str, row: int, max_utf16: int) -> str:
    if type(value) not in (str, int) or isinstance(value, bool):
        raise DraftImportError("Textwert fehlt oder ist ungueltig.", row=row, field=field)
    text = str(value).strip()
    if (not text or len(text.encode("utf-16-le", "surrogatepass")) // 2 > max_utf16
        or any((ord(c) < 32 and c not in "\n\t") or 0xD800 <= ord(c) <= 0xDFFF for c in text)):
        raise DraftImportError("Textwert ungueltig oder zu lang.", row=row, field=field)
    return text


def _draft(values: list[object], mapping: dict[str, int], row: int) -> ProductDraft:
    def get(field: str) -> object:
        i = mapping.get(field)
        return values[i] if i is not None and i < len(values) else None

    sku = _text(get("sku"), "sku", row, 80)
    if not _SKU.fullmatch(sku):
        raise DraftImportError("Artikelnummer nur Buchstaben, Ziffern, Punkt, _ und -.", row=row, field="sku")
    title = _text(get("title"), "title", row, 65)
    description = _text(get("description"), "description", row, 4000)
    category = tuple(part.strip() for part in _text(get("category"), "category", row, 750).split(">"))
    if not 2 <= len(category) <= 6 or any(not name or len(name) > 120 for name in category):
        raise DraftImportError("Kategorie braucht 2 bis 6 nichtleere Teile.", row=row, field="category")
    raw_price = get("price_eur")
    if type(raw_price) is int:
        price = raw_price
    elif type(raw_price) is float and math.isfinite(raw_price) and raw_price.is_integer():
        price = int(raw_price)
    elif type(raw_price) is str and re.fullmatch(r"[0-9]{1,8}(?:[.,]00)?", raw_price.strip()):
        price = int(raw_price.strip().split(".")[0].split(",")[0])
    else:
        raise DraftImportError("Festpreis muss volle Euro enthalten.", row=row, field="price_eur")
    if not 1 <= price <= 99_999_999:
        raise DraftImportError("Festpreis ausserhalb des gueltigen Bereichs.", row=row, field="price_eur")
    raw_images = get("images")
    if raw_images in (None, ""):
        images: tuple[str, ...] = ()
    else:
        if not isinstance(raw_images, str):
            raise DraftImportError("Bilddateinamen muessen Text sein.", row=row, field="images")
        items = [part.strip() for part in raw_images.split(";")]
        if not 1 <= len(items) <= 10 or len(set(x.casefold() for x in items)) != len(items):
            raise DraftImportError("Maximal 10 unterschiedliche Bilddateinamen.", row=row, field="images")
        if any(not _IMG.fullmatch(x) or x in (".", "..") for x in items):
            raise DraftImportError("Bilder nur Dateinamen, keine Pfade.", row=row, field="images")
        images = tuple(items)
    return ProductDraft(sku, title, description, category, price, images)


def _read_rows(path: Path, content: bytes) -> tuple[list[list[object]], str]:
    if path.suffix.casefold() == ".csv":
        try:
            decoded = content.decode("utf-8-sig")
            with io.StringIO(decoded, newline="") as stream:
                reader = csv.reader(stream, delimiter=",", strict=True)
                values = []
                for index, row in enumerate(reader):
                    if index > MAX_ROWS:
                        raise DraftImportError("Maximal 5000 Artikel pro Import.")
                    values.append(list(row))
                return values, "csv"
        except (UnicodeError, csv.Error):
            raise DraftImportError("CSV muss UTF-8 mit Kommatrennung sein.", field="file") from None
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if (len(members) > 500 or sum(x.file_size for x in members) > _MAX_DECOMPRESSED
                or any(x.file_size > _MAX_DECOMPRESSED for x in members)):
                raise DraftImportError("XLSX-Datei ist zu gross oder ungueltig.", field="file")
            if len({x.filename for x in members}) != len(members):
                raise DraftImportError("XLSX enthaelt doppelte Dateieintraege.", field="file")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.active
            values = []
            for n, cells in enumerate(sheet.iter_rows(max_row=MAX_ROWS+2), start=1):
                if n > MAX_ROWS + 1:
                    raise DraftImportError("Maximal 5000 Artikel pro Import.")
                if any(cell.data_type in ("f", "e") for cell in cells):
                    raise DraftImportError("Formeln und Excel-Fehler sind nicht erlaubt.", row=n)
                values.append([cell.value for cell in cells])
        finally:
            workbook.close()
        return values, "xlsx"
    except DraftImportError:
        raise
    except Exception:  # noqa: BLE001 - no corrupt Excel details escape into customer UI
        raise DraftImportError("XLSX konnte nicht sicher gelesen werden.", field="file") from None


def parse_drafts(path: Path) -> tuple[list[ProductDraft], dict[str, str], str, str]:
    contents, digest = _read_file(path)
    raw, file_format = _read_rows(path, contents)
    if not raw:
        raise DraftImportError("Datei enthaelt keine Kopfzeile.", row=1)
    mapping = _column_map(raw[0])
    result: list[ProductDraft] = []
    seen: set[str] = set()
    issues: list[dict[str, object]] = []
    for row_no, values in enumerate(raw[1:], start=2):
        if not any(x is not None and str(x).strip() for x in values):
            continue
        try:
            if len(values) > len(raw[0]):
                raise DraftImportError("Mehr Datenzellen als Spalten in dieser Zeile.", row=row_no)
            draft = _draft(values, mapping, row_no)
            if draft.sku in seen:
                raise DraftImportError("Doppelte Artikelnummer in Importdatei.", row=row_no, field="sku")
            seen.add(draft.sku)
            result.append(draft)
        except DraftImportError as exc:
            if len(issues) < 25:
                issues.append(exc.as_issue())
    if issues:
        raise DraftBatchError(issues)
    if not result:
        raise DraftImportError("Datei enthaelt keine Artikel.", row=2)
    return result, {field: str(raw[0][pos]) for field, pos in mapping.items()}, digest, file_format


def _has_table(conn: sqlite3.Connection) -> bool:
    row = conn.execute("SELECT type FROM sqlite_master WHERE name='portfolio_drafts'").fetchone()
    if row is None:
        return False
    if row[0] != "table":
        raise sqlite3.DatabaseError("draft schema is invalid")
    return True


def _stored(conn: sqlite3.Connection, sku: str) -> tuple[object, ...] | None:
    row = conn.execute(
        "SELECT title,description,category_json,price_eur,images_json,status "
        "FROM portfolio_drafts WHERE sku = ? COLLATE BINARY", (sku,)
    ).fetchone()
    if row is None:
        return None
    if row[5] != "draft":
        raise sqlite3.DatabaseError("draft state is invalid")
    return tuple(row[:5])


def preview_import(store: SnapshotStore, path: Path) -> dict[str, object]:
    rows, mapping, digest, file_format = parse_drafts(path)
    count = {"new": 0, "unchanged": 0, "conflict": 0}
    with closing(store._connect()) as conn:
        exists = _has_table(conn)
        for item in rows:
            old = _stored(conn, item.sku) if exists else None
            key = "new" if old is None else ("unchanged" if old == item.fields() else "conflict")
            count[key] += 1
    return {"status": "preview", "filename": path.name, "format": file_format,
            "sha256": digest, "total": len(rows), "mapping": mapping,
            **count, "samples": [r.as_dict() for r in rows[:6]],
            "can_import": count["conflict"] == 0}


def import_drafts(store: SnapshotStore, path: Path, expected_sha256: str) -> dict[str, object]:
    if type(expected_sha256) is not str or not _HEX.fullmatch(expected_sha256):
        raise DraftImportError("Ungueltiger Datei-Pruefwert.", field="file")
    rows, _, digest, _ = parse_drafts(path)
    if digest != expected_sha256:
        raise DraftImportError("Datei wurde nach der Vorschau geaendert.", field="file")
    now = datetime.now(timezone.utc).isoformat()
    inserted = 0
    identical = 0
    with closing(store._connect()) as conn, conn:
        conn.execute("BEGIN IMMEDIATE")
        conn.execute(_DDL)
        for item in rows:
            old = _stored(conn, item.sku)
            if old is not None:
                if old != item.fields():
                    raise DraftImportError("Artikelnummer existiert mit anderem Inhalt.", field="sku")
                identical += 1
                continue
            conn.execute(
                "INSERT INTO portfolio_drafts "
                "(sku,title,description,category_json,price_eur,images_json,"
                "source_sha256,source_name,created_at,status) VALUES (?,?,?,?,?,?,?,?,?,'draft')",
                (item.sku, *item.fields(), digest, path.name, now),
            )
            inserted += 1
    return {"status": "imported", "total": len(rows), "inserted": inserted,
            "unchanged": identical, "published": 0}


def drafts_page(store: SnapshotStore, *, after: str = "", search: str = "", limit: int = 50) -> dict[str, object]:
    if (type(limit) is not int or not 1 <= limit <= 50
        or type(after) is not str or len(after) > 80
        or (after and not _SKU.fullmatch(after))
        or type(search) is not str or len(search) > 120
        or any(ord(c) < 32 for c in search)):
        raise ValueError("invalid draft catalog query")
    with closing(store._connect()) as conn:
        if not _has_table(conn):
            return {"drafts": [], "has_more": False, "next_after": None, "total": 0}
        conn.create_function("unicode_casefold", 1, lambda text: text.casefold(), deterministic=True)
        where = "sku > ? COLLATE BINARY"
        args: list[object] = [after]
        if search:
            where += " AND (instr(unicode_casefold(sku), ?) > 0 OR instr(unicode_casefold(title), ?) > 0)"
            args.extend((search.casefold(), search.casefold()))
        total = conn.execute(
            "SELECT COUNT(*) FROM portfolio_drafts" +
            (" WHERE instr(unicode_casefold(sku), ?) > 0 OR instr(unicode_casefold(title), ?) > 0" if search else ""),
            ((search.casefold(), search.casefold()) if search else ()),
        ).fetchone()[0]
        records = conn.execute(
            "SELECT sku,title,description,category_json,price_eur,images_json,status "
            "FROM portfolio_drafts WHERE " + where + " ORDER BY sku COLLATE BINARY LIMIT ?",
            (*args, limit+1),
        ).fetchall()
    visible = records[:limit]
    data = []
    for rec in visible:
        try:
            category, images = json.loads(rec[3]), json.loads(rec[5])
            if rec[6] != "draft" or not isinstance(category, list) or not isinstance(images, list):
                raise ValueError("invalid stored draft")
            data.append({"sku": rec[0], "title": rec[1], "description": rec[2],
                         "category": " > ".join(category), "price_eur": rec[4],
                         "images": images, "status": "draft", "published": False})
        except (TypeError, ValueError):
            raise sqlite3.DatabaseError("invalid draft data") from None
    return {"drafts": data, "has_more": len(records) > limit,
            "next_after": visible[-1][0] if len(records) > limit else None, "total": total}


def _excel_csv_text(value: object) -> object:
    if isinstance(value, str):
        leading = value.lstrip()
        if leading.startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
            return "'" + value
    return value


def export_drafts(store: SnapshotStore, target: Path) -> dict[str, object]:
    """Write a complete, Excel-safe CSV through a private stage, create-only.

    The target is published as a hardlink only after the full CSV has been
    flushed; no existing customer CSV can be replaced and partial exports are
    removed on failures.
    """
    if target.suffix.casefold() != ".csv" or not target.is_absolute():
        raise DraftImportError("Exportziel muss eine absolute CSV-Datei sein.")
    with closing(store._connect()) as conn:
        if not _has_table(conn):
            records = []
        else:
            records = conn.execute(
                "SELECT sku,title,description,category_json,price_eur,images_json "
                "FROM portfolio_drafts ORDER BY sku COLLATE BINARY"
            ).fetchall()
    stage = target.with_name("." + target.name + "." + secrets.token_hex(10) + ".part")
    try:
        with stage.open("x", encoding="utf-8-sig", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(_FIELDS)
            for row in records:
                category = " > ".join(json.loads(row[3]))
                images = ";".join(json.loads(row[5]))
                writer.writerow(tuple(_excel_csv_text(x) for x in
                                      (row[0], row[1], row[2], category, row[4], images)))
            file.flush()
            os.fsync(file.fileno())
        os.link(stage, target)
    except (OSError, ValueError, TypeError):
        raise DraftImportError("Entwuerfe konnten nicht exportiert werden.") from None
    finally:
        try:
            stage.unlink(missing_ok=True)
        except OSError:
            pass
    return {"status": "exported", "rows": len(records), "filename": target.name}
